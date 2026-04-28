from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# Set margins to 1 inch (2.54 cm)
for section in doc.sections:
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

# Add page numbers in footer
for section in doc.sections:
    footer = section.footer
    footer_para = footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_para.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)

def set_font(run, size=11, bold=False, color=None):
    run.font.name = 'Calibri'
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)

def add_heading(doc, text, level):
    para = doc.add_heading(text, level=level)
    for run in para.runs:
        run.font.name = 'Calibri'
    return para

def add_paragraph(doc, text='', bold=False, size=11):
    para = doc.add_paragraph()
    if text:
        run = para.add_run(text)
        set_font(run, size=size, bold=bold)
    return para

def add_bullet(doc, text, level=0):
    para = doc.add_paragraph(style='List Bullet')
    run = para.add_run(text)
    set_font(run, size=11)
    return para

# ─── TITLE PAGE ────────────────────────────────────────────────────────────────

title_para = doc.add_paragraph()
title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title_para.add_run("Hamna's Part – User Management, Forgot Password & Staff Profile")
set_font(run, size=20, bold=True, color=(0, 70, 127))

doc.add_paragraph()

subj_para = doc.add_paragraph()
subj_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subj_para.add_run("IT2021: AIML Project – Assignment 03 Progress II")
set_font(run, size=14, bold=True)

doc.add_paragraph()

inst_para = doc.add_paragraph()
inst_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = inst_para.add_run("Sri Lanka Institute of Information Technology")
set_font(run, size=13, bold=True)

dept_para = doc.add_paragraph()
dept_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = dept_para.add_run("Department of IT, Faculty of Computing")
set_font(run, size=12)

yr_para = doc.add_paragraph()
yr_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = yr_para.add_run("2nd Year, Semester 2, 2026")
set_font(run, size=12)

doc.add_paragraph()

student_para = doc.add_paragraph()
student_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = student_para.add_run("Student: Hamna")
set_font(run, size=12, bold=True)

system_para = doc.add_paragraph()
system_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = system_para.add_run("System: Invigo FreshGuard – Perishable Goods Inventory Management System")
set_font(run, size=11)

doc.add_page_break()

# ─── SECTION 1 ─────────────────────────────────────────────────────────────────

add_heading(doc, "Section 1: Responsible Components (~75% Implemented)", 1)

p = add_paragraph(doc, "This section details the three frontend modules developed and implemented by Hamna for the Invigo FreshGuard system. Each module has been implemented to approximately 75% completion and covers the core functional requirements.")
for run in p.runs:
    set_font(run, size=11)

# 1.1 User Management
add_heading(doc, "1.1 User Management (/admin/users)", 2)

p = add_paragraph(doc, "The User Management module provides administrators with full control over staff user accounts, security monitoring, and role-based access configuration.")
for run in p.runs:
    set_font(run, size=11)

add_heading(doc, "Feature 1: User Table with Filters & Sorting", 3)
p = add_paragraph(doc, "The main view displays all staff users in a structured table. The module supports text search by name or username, a role filter dropdown, and a date-of-joining range filter using dojFrom and dojTo parameters. A validation rule prevents the end date from being set before the start date. Each column (Name, Username, Role, Date of Joining) supports click-to-sort with ascending/descending toggle, and an active sort indicator provides immediate visual feedback to the user.")
for run in p.runs:
    set_font(run, size=11)

add_heading(doc, "Feature 2: Create User", 3)
p = add_paragraph(doc, "Administrators can add a new staff member via a modal dialog. The form includes the following fields: Full Name, Username, Password (with show/hide toggle), Date of Joining, Role (populated dynamically from the backend), and an optional Email field. Validation rules enforce that all required fields are present, the full name is between 2 and 50 characters, the password is between 6 and 30 characters, and the date of joining is not a future date. On successful submission, the new user is created in the backend and appended to the table without a page reload.")
for run in p.runs:
    set_font(run, size=11)

add_heading(doc, "Feature 3: Edit User", 3)
p = add_paragraph(doc, "Clicking the Edit button on any user row opens a pre-populated dialog. The administrator can update the name, username, date of joining, role, and optionally set a new password. All required fields are validated before submission. On success, the backend is updated and the corresponding table row is updated in place, providing a seamless editing experience.")
for run in p.runs:
    set_font(run, size=11)

add_heading(doc, "Feature 4: Delete User", 3)
p = add_paragraph(doc, "Clicking the trash icon on a non-admin user row triggers a confirmation dialog before deletion. On confirmation, the user is removed from the backend and the corresponding row is removed from the table. Admin accounts are protected from deletion — the delete button is hidden or disabled for any row belonging to an admin account.")
for run in p.runs:
    set_font(run, size=11)

add_heading(doc, "Feature 5: Account Lock / Unlock", 3)
p = add_paragraph(doc, "The system tracks failed login attempts per user. Accounts that have been locked after repeated failed attempts display a 'Locked' badge in the user table. Administrators can unlock individual accounts with a single click. Additionally, the Admin Dashboard displays a live security alert banner whenever any staff accounts are locked, including a direct link to the User Control page.")
for run in p.runs:
    set_font(run, size=11)

add_heading(doc, "Feature 6: Security Activity Tab", 3)
p = add_paragraph(doc, "A dedicated sub-tab inside User Control provides a comprehensive security audit view. It includes:")
for run in p.runs:
    set_font(run, size=11)
add_bullet(doc, "Login History Table: displays all login events with user identity, date/time, and a colour-coded status badge (Success = green, Failed = red).")
add_bullet(doc, "Search and filter capability by name/username and status (All, Success, Failed, Locked, Role Mismatch).")
add_bullet(doc, "Recent Changes Card: shows the last 5 login events on a visual timeline.")
add_bullet(doc, "Failed Login Attempts & Locks Card: lists all staff members with failed attempts greater than zero, each with an inline Unlock button.")

add_heading(doc, "Feature 7: Role Permissions Tab", 3)
p = add_paragraph(doc, "A separate sub-tab provides full role management capabilities. It includes:")
for run in p.runs:
    set_font(run, size=11)
add_bullet(doc, "A table displaying all roles with role name, description, assigned user count, and last modified date.")
add_bullet(doc, "Clicking a role opens a side panel with 11 granular permission checkboxes (e.g., Sales Recording, Edit Products, User Management, etc.). A 'Save Changes' button updates role permissions in the backend.")
add_bullet(doc, "A 'Create New Role' dialog requiring a role name (mandatory) and optional description.")
add_bullet(doc, "Delete role functionality with a confirmation dialog.")
add_bullet(doc, "Search and filter roles by name.")

# 1.2 Forgot Password
add_heading(doc, "1.2 Forgot Password (/forgot-password)", 2)

p = add_paragraph(doc, "The Forgot Password module provides a guided multi-step password recovery flow, designed for clarity and ease of use across desktop and mobile devices.")
for run in p.runs:
    set_font(run, size=11)

add_heading(doc, "Feature 1: 4-Step Animated Flow", 3)
p = add_paragraph(doc, "The recovery process is structured as a multi-step wizard with animated slide transitions (left/right based on navigation direction). A visual step indicator on the left panel displays the four stages: Email, OTP, Password, and Done. A progress bar at the top of the form card provides an additional visual cue of the user's current position in the flow.")
for run in p.runs:
    set_font(run, size=11)

add_heading(doc, "Feature 2: Step 1 – Email Entry", 3)
p = add_paragraph(doc, "The user enters their registered email address. Real-time inline validation applies a regex check — the field border turns red with an error message if the format is invalid, and turns green when the format is valid. The submit button remains disabled while the email is in an invalid format. On submission, the frontend calls the /api/forgot-password backend endpoint.")
for run in p.runs:
    set_font(run, size=11)

add_heading(doc, "Feature 3: Step 2 – OTP Verification", 3)
p = add_paragraph(doc, "The OTP entry screen presents 6 individual digit input boxes. Typing a digit automatically moves focus to the next box; pressing Backspace moves focus to the previous box. Clipboard paste support allows users to paste a 6-digit code and have all boxes filled automatically — a significant usability improvement for mobile users. A 60-second countdown timer locks the Resend OTP button; once expired, the user may request a new OTP which resets the timer. Submission calls /api/verify-otp.")
for run in p.runs:
    set_font(run, size=11)

add_heading(doc, "Feature 4: Step 3 – New Password", 3)
p = add_paragraph(doc, "Both password and confirm password fields include show/hide eye toggles. A real-time 4-segment password strength bar provides feedback (Too Short / Fair / Strong). The password field shows a red border and error text if the password is fewer than 6 characters. The confirm field shows a red border if it does not match the password, and a green confirmation message when both match and the minimum length is met. The submit button is disabled until all conditions are satisfied. On submission, the frontend calls /api/reset-password.")
for run in p.runs:
    set_font(run, size=11)

add_heading(doc, "Feature 5: Step 4 – Success Screen", 3)
p = add_paragraph(doc, "The final step displays an animated success icon with a spring animation, a confirmation message, and a 'Go to Login' button that navigates the user to /login.")
for run in p.runs:
    set_font(run, size=11)

add_heading(doc, "Feature 6: Back Navigation", 3)
p = add_paragraph(doc, "A 'Change email' link on the OTP step allows the user to return to the email entry step. A 'Back to Login' link on the email step navigates the user back to /login.")
for run in p.runs:
    set_font(run, size=11)

# 1.3 Staff Profile
add_heading(doc, "1.3 Staff Profile (/staff/profile)", 2)

p = add_paragraph(doc, "The Staff Profile module allows logged-in staff members to view their profile information, update their display name, and change their password.")
for run in p.runs:
    set_font(run, size=11)

add_heading(doc, "Feature 1: Profile Card", 3)
p = add_paragraph(doc, "The profile card displays the logged-in staff member's initials as an avatar, along with their full name, username, role (with normalized capitalization), and date of joining formatted as 'Month Day, Year'.")
for run in p.runs:
    set_font(run, size=11)

add_heading(doc, "Feature 2: Edit Display Name", 3)
p = add_paragraph(doc, "A dedicated form allows the staff member to update their display name. Validation prevents empty names from being submitted. If the submitted name is identical to the current name, a 'No changes' toast notification is shown instead of making an API call. On a successful update, the session data in localStorage is updated so the sidebar name refreshes immediately, providing instant visual confirmation.")
for run in p.runs:
    set_font(run, size=11)

add_heading(doc, "Feature 3: Change Password", 3)
p = add_paragraph(doc, "A password change form includes three fields — current password, new password, and confirm password — each equipped with show/hide eye toggles. Validation is performed in strict priority order: current password required, new password required, new password must be at least 8 characters, confirm must match new password, and new password must differ from the current password. Inline error text is shown for each violation. On success, a toast notification confirms the change.")
for run in p.runs:
    set_font(run, size=11)

add_heading(doc, "Feature 4: Loading & Error States", 3)
p = add_paragraph(doc, "A full-page spinner is displayed while the profile data is being fetched from the backend. If the backend is unreachable, a styled error card with a 'Retry Connection' button is presented. Users who are not authenticated (no valid session in localStorage) are automatically redirected to /login.")
for run in p.runs:
    set_font(run, size=11)

doc.add_page_break()

# ─── SECTION 2 ─────────────────────────────────────────────────────────────────

add_heading(doc, "Section 2: User Experience – Task Flow, Navigation & Usability", 1)

add_heading(doc, "2.1 User Management", 2)
p = add_paragraph(doc, "The User Management module uses a table-and-filter-bar layout that keeps all tools visible on a single screen without requiring additional navigation. Sort arrows and an active sort indicator provide immediate visual feedback when a column is sorted. The locked account security alert banner on the Admin Dashboard drives administrator attention to urgent account security issues and provides a direct navigation link, reducing the time required to respond to a security event.")
for run in p.runs:
    set_font(run, size=11)

add_heading(doc, "2.2 Forgot Password", 2)
p = add_paragraph(doc, "The 4-step wizard structure prevents user confusion by presenting only one task at a time. The persistent step indicator on the left panel gives users a clear understanding of their progress through the recovery flow at all times. Error messages appear inline, immediately adjacent to the offending field, so users do not need to search for the source of validation failures. OTP clipboard paste support eliminates friction for mobile users who receive OTP codes via SMS or email, as they can paste directly without manual digit entry.")
for run in p.runs:
    set_font(run, size=11)

add_heading(doc, "2.3 Staff Profile", 2)
p = add_paragraph(doc, "The Staff Profile page uses a clean two-section layout, with profile information displayed above and the password change form below. Feedback is delivered via both inline messages and toast notifications, ensuring users receive confirmation regardless of where they are looking on the screen. The loading state prevents a blank screen experience during data fetch. The immediate sidebar name update after a successful profile save provides instant visual confirmation that the change has taken effect.")
for run in p.runs:
    set_font(run, size=11)

doc.add_page_break()

# ─── SECTION 3 ─────────────────────────────────────────────────────────────────

add_heading(doc, "Section 3: UI Consistency & Standards", 1)

p = add_paragraph(doc, "All components developed by Hamna adhere to the established visual language and design standards of the Invigo FreshGuard system.")
for run in p.runs:
    set_font(run, size=11)

bullets3 = [
    "The admin panel uses a dark sidebar (#0F172A) with white text, consistent across all admin-facing pages including /admin/users.",
    "The user table uses the same font-black uppercase 10px tracking-widest header style as all other admin data tables in the system.",
    "All modal dialogs follow the same rounded-2xl structure with a consistent label → input → error text layout.",
    "Status badges use a consistent colour coding convention: green for Success/Active, red for Failed/Locked, and orange for Warning states.",
    "The Forgot Password page mirrors the Login page layout — split left panel and right card, the same blob background decorations, and the same brand colour palette — ensuring visual continuity.",
    "The Staff Profile page uses the warm staff theme (#F9F5EC background, #4E342E accent), consistent with all other /staff/* pages in the system.",
    "All password fields across all three modules use the same show/hide Eye/EyeOff toggle pattern, ensuring a consistent interaction model.",
    "All forms use the same red-50 background error box pattern with border-red-100 and an AlertTriangle icon for form-level error messages.",
]
for b in bullets3:
    add_bullet(doc, b)

doc.add_page_break()

# ─── SECTION 4 ─────────────────────────────────────────────────────────────────

add_heading(doc, "Section 4: Input Validation & Error Handling", 1)

p = add_paragraph(doc, "All three modules implement comprehensive client-side input validation to ensure data integrity and provide clear, actionable error feedback to users.")
for run in p.runs:
    set_font(run, size=11)

add_heading(doc, "4.1 Create User Form", 2)
validations_create = [
    ('All fields required (username, password, name, DOJ)', '"Please fill in all fields to continue."'),
    ('Full name: 2–50 characters', '"Full name must be between 2 and 50 characters."'),
    ('Password: 6–30 characters', '"Password must be between 6 and 30 characters."'),
    ('Date of Joining: cannot be a future date', '"Joining date cannot be a future date."'),
    ('DOJ date range filter: end date cannot be before start date', 'Inline error shown on filter bar'),
]
for rule, error in validations_create:
    para = doc.add_paragraph(style='List Bullet')
    run1 = para.add_run(rule + ' — ')
    set_font(run1, size=11, bold=True)
    run2 = para.add_run('Error: ' + error)
    set_font(run2, size=11)

add_heading(doc, "4.2 Edit User Form", 2)
p = add_paragraph(doc, 'Username, name, and DOJ are required — Error: "Please fill in all required fields."')
for run in p.runs:
    set_font(run, size=11)

add_heading(doc, "4.3 Role Management", 2)
p = add_paragraph(doc, 'Create Role: role name is required — Error: "Role name is required."')
for run in p.runs:
    set_font(run, size=11)

add_heading(doc, "4.4 Forgot Password – Step 1 (Email)", 2)
fp1_bullets = [
    'Real-time regex validation using /^[^\\s@]+@[^\\s@]+\\.[^\\s@]+$/',
    'Invalid email: red border + inline "Invalid email address" text',
    'Valid email: green border + "Valid email address" confirmation',
    'Submit button is disabled while the email format is invalid',
]
for b in fp1_bullets:
    add_bullet(doc, b)

add_heading(doc, "4.5 Forgot Password – Step 2 (OTP)", 2)
fp2_bullets = [
    'OTP must be exactly 6 digits — Error: "Please enter all 6 digits."',
    'Submit button is disabled until all 6 input boxes are filled',
    'Only digit characters are accepted (regex /^\\d?$/ filter applied per input box)',
]
for b in fp2_bullets:
    add_bullet(doc, b)

add_heading(doc, "4.6 Forgot Password – Step 3 (New Password)", 2)
fp3_bullets = [
    'Minimum 6 characters — Error: "Password must be at least 6 characters."',
    'Passwords must match — Error: "Passwords do not match"',
    'Submit button is disabled until length >= 6 AND both passwords match',
    'Real-time visual strength bar provides feedback as the user types',
]
for b in fp3_bullets:
    add_bullet(doc, b)

add_heading(doc, "4.7 Staff Profile – Change Password", 2)
sp_bullets = [
    'Current password required — Error: "Please enter your current password."',
    'New password required — Error: "Please enter a new password."',
    'New password minimum 8 characters — Error: "New password must be at least 8 characters."',
    'Confirm password must match new password — Error: "New passwords do not match."',
    'New password must differ from current — Error: "New password must be different from your current password."',
    'Validations run in strict priority order to surface the most actionable error first',
]
for b in sp_bullets:
    add_bullet(doc, b)

add_heading(doc, "4.8 Staff Profile – Edit Name", 2)
sp_name_bullets = [
    'Empty name is not accepted — guarded with a no-op check before API call',
    'Unchanged name detected — shows "No changes" toast instead of making an API call',
]
for b in sp_name_bullets:
    add_bullet(doc, b)

add_heading(doc, "4.9 Error Display Patterns", 2)
error_patterns = [
    'Form-level errors: red-50 background box with border-red-100 and an AlertTriangle icon',
    'Field-level errors: inline <p> element in red-500 colour displayed below the input field',
    'Toast notifications: destructive variant (red background) for errors, default variant for success',
]
for b in error_patterns:
    add_bullet(doc, b)

doc.add_page_break()

# ─── SECTION 5 ─────────────────────────────────────────────────────────────────

add_heading(doc, "Section 5: Testing of Responsible Components", 1)

add_heading(doc, "5.1 Frontend Unit Tests (Vitest)", 2)

add_heading(doc, "userManagement.test.js — 13 Test Cases", 3)
um_tests = [
    'Returns error when all fields are empty',
    'Returns error for missing username',
    'Returns error for missing password',
    'Returns error for missing name',
    'Returns error for missing date of joining',
    'Boundary: password exactly 5 characters fails validation',
    'Boundary: password exactly 6 characters passes validation',
    'Returns null (no error) for a fully valid input set',
    'Edit User: error when username is missing',
    'Edit User: error when name is missing',
    'Edit User: error when date of joining is missing',
    'Edit User: returns null when all required fields are valid',
    'DOJ range filter: returns error when end date is before start date',
]
for t in um_tests:
    add_bullet(doc, t)

add_heading(doc, "staffProfile.test.js — 9 Test Cases", 3)
sp_tests = [
    'Returns error when current password is missing',
    'Returns error when new password is missing',
    'Boundary: new password with 7 characters fails the 8-character minimum',
    'Boundary: new password with 8 characters passes the minimum length check',
    'Returns error when confirm password does not match new password',
    'Returns error when new password is identical to current password',
    'Returns null (no error) for a fully valid input set',
    'Priority order: verifies the correct error is returned when multiple validation rules are violated simultaneously',
    'No-op guard: unchanged display name triggers "No changes" toast rather than API call',
]
for t in sp_tests:
    add_bullet(doc, t)

add_heading(doc, "authGuard.test.js — 12 Test Cases", 3)
p = add_paragraph(doc, "Route protection tests covering the /admin/users route, which requires ADMIN role enforcement. Full details are covered under the authentication section; relevant to User Management because unauthenticated or insufficiently privileged access must be blocked at the route level.")
for run in p.runs:
    set_font(run, size=11)

add_heading(doc, "5.2 Manual Test Cases", 2)

p = add_paragraph(doc, "The following manual test cases were conducted and are demonstrable:")
for run in p.runs:
    set_font(run, size=11)

doc.add_paragraph()

# Build test case table
headers = ['#', 'Test Case', 'Input', 'Expected', 'Result']
test_cases = [
    ('1', 'Create user with all valid fields', 'Valid username, password, name, DOJ', 'User created and added to table', 'Pass'),
    ('2', 'Create user with empty fields', 'All fields blank', 'Error: fill all fields', 'Pass'),
    ('3', 'Create user with short password', 'Password = "abc"', 'Error: 6–30 characters', 'Pass'),
    ('4', 'Create user with future DOJ', 'DOJ = tomorrow', 'Error: cannot be future', 'Pass'),
    ('5', 'Edit user with missing name', 'Name field blank', 'Error: required fields', 'Pass'),
    ('6', 'Delete admin user', 'Admin row selected', 'Delete button hidden', 'Pass'),
    ('7', 'Unlock locked account', 'Locked user account', 'Account unlocked, badge removed', 'Pass'),
    ('8', 'FP – invalid email format', '"notanemail"', 'Red border + inline error', 'Pass'),
    ('9', 'FP – OTP with wrong code', 'Incorrect 6-digit code', 'Error toast from backend', 'Pass'),
    ('10', 'FP – password too short', '4 character password', 'Submit disabled + red border', 'Pass'),
    ('11', "FP – passwords don't match", 'pw1 not equal to pw2', 'Red border + error text', 'Pass'),
    ('12', 'FP – paste OTP', 'Paste 6-digit code', 'All boxes filled automatically', 'Pass'),
    ('13', 'FP – resend OTP before 60s', 'Click Resend during countdown', 'Button disabled, shows timer', 'Pass'),
    ('14', 'Profile – change pw same as current', 'Same value for both fields', 'Error: must be different', 'Pass'),
    ('15', 'Profile – new pw < 8 chars', '5 character password', 'Error: minimum 8 characters', 'Pass'),
    ('16', 'Profile – edit name unchanged', 'Submit same name', 'Toast: No changes', 'Pass'),
    ('17', 'Profile – unauthenticated access', 'No localStorage session', 'Redirected to /login', 'Pass'),
    ('18', 'Login history – filter by Failed', 'Select "Failed" in dropdown', 'Only failed logins shown', 'Pass'),
    ('19', 'Role create with empty name', 'Blank role name field', 'Error: name is required', 'Pass'),
    ('20', 'DOJ range filter: end before start', 'dojTo < dojFrom', 'Error: end cannot be before start', 'Pass'),
]

col_widths = [Cm(0.8), Cm(4.5), Cm(3.5), Cm(4.0), Cm(1.8)]
table = doc.add_table(rows=1, cols=5)
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

# Header row
hdr_cells = table.rows[0].cells
for i, (cell, header, width) in enumerate(zip(hdr_cells, headers, col_widths)):
    cell.width = width
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(header)
    run.font.name = 'Calibri'
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), '1F4E79')
    tcPr.append(shd)

# Data rows
for row_data in test_cases:
    row_cells = table.add_row().cells
    for i, (cell, val, width) in enumerate(zip(row_cells, row_data, col_widths)):
        cell.width = width
        para = cell.paragraphs[0]
        if i == 0:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(val)
        run.font.name = 'Calibri'
        run.font.size = Pt(10)
        if i == 4:  # Result column
            run.font.color.rgb = RGBColor(0, 128, 0)
            run.font.bold = True

doc.add_page_break()

# ─── SECTION 6 ─────────────────────────────────────────────────────────────────

add_heading(doc, "Section 6: Backend Components (Brief)", 1)

p = add_paragraph(doc, "The following backend components support the three modules described in this document. These components were developed in Java (Spring Boot) and interact with the frontend via REST API endpoints.")
for run in p.runs:
    set_font(run, size=11)

backend_items = [
    ('AuthController.java', 'Handles the /api/forgot-password, /api/verify-otp, and /api/reset-password endpoints for the Forgot Password flow.'),
    ('PasswordResetController.java', 'Manages OTP token generation, expiry validation, and verification logic.'),
    ('OtpToken (entity)', 'Stores the email address, 6-digit OTP code, and expiry timestamp for each active password reset request.'),
    ('AdminUserController.java', 'Handles full CRUD operations for staff user accounts and the account unlock functionality.'),
    ('StaffProfileController.java', 'Handles GET and PUT operations for staff profile data, including display name update and password change.'),
    ('LoginHistoryController.java / LoginHistoryService.java', 'Records every login attempt with its associated status, providing data for the Security Activity tab.'),
    ('RoleController.java', 'Handles CRUD operations for roles and their associated granular permissions.'),
    ('SecurityConfig.java', 'Configures BCrypt password hashing, stateless session management, and CORS settings for the application.'),
]

for component, description in backend_items:
    para = doc.add_paragraph(style='List Bullet')
    run1 = para.add_run(component + ': ')
    set_font(run1, size=11, bold=True)
    run2 = para.add_run(description)
    set_font(run2, size=11)

# Save
output_path = r"c:\Users\ASUS\Downloads\invigo-freshguard-main\invigo-freshguard-main\Hamna's Part.docx"
doc.save(output_path)
print(f"Saved: {output_path}")