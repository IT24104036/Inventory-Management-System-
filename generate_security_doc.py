from docx import Document
import os

def create_document():
    # Initialize document
    doc = Document()

    # Title
    doc.add_heading('Security Testing Justification: Login & OTP Pages', 0)

    # Section 1
    doc.add_heading('1. Why do we need SQL Injection testing on the Login Page?', level=2)
    doc.add_paragraph(
        "The login page takes user input (email/username and password) and sends it directly to our database to verify the user. "
        "If we don’t test for SQL injection and sanitize those inputs, an attacker could enter malicious SQL commands (like ' OR 1=1 --) in the username field. "
        "This would trick the database into validating the login as true, allowing the attacker to bypass authentication entirely and gain unauthorized access to the system"
        "—perhaps even as an administrator—without ever knowing a real password. Testing this ensures our database is safe from manipulation."
    )

    # Section 2
    doc.add_heading('2. Why do we need Brute-force testing on the Login Page?', level=2)
    doc.add_paragraph(
        "We need to test for brute-force attacks on the login page because attackers use automated scripts to guess thousands of passwords per second. "
        "If we don't test for this and implement protections (like rate-limiting or account lockouts after too many failed attempts), an attacker will eventually guess a weak password. "
        "Security testing proves that our system will block those rapid-fire malicious attempts, protecting user accounts from being compromised."
    )

    # Section 3
    doc.add_heading('3. Why do we need Brute-force testing on the OTP Page?', level=2)
    doc.add_paragraph(
        "An OTP (One-Time Password) is usually only 4 or 6 digits long. A 4-digit pin only has 10,000 possible combinations. "
        "If an attacker's script can submit 100 guesses a second, they can crack the OTP and bypass our Two-Factor Authentication in less than two minutes. "
        "Security testing on the OTP page ensures that our validation logic works—specifically, that the OTP expires after a short time, and that the system locks the user out after a few failed attempts (rate limiting), making brute-forcing mathematically impossible."
    )

    # Section 4
    doc.add_heading('Alignment with Rubric (Progress II)', level=2)
    doc.add_paragraph(
        "Implementing these tests directly covers the 'Input Validation & Error Handling' and 'Basic Security Testing' requirements of our Progress II milestone. "
        "As the developer handling this part of the project, it was my responsibility to ensure that the core entry points of our application are not just functional, but secure against standard web vulnerabilities."
    )

    # Save document
    file_name = 'Security_Testing_Justification.docx'
    doc.save(file_name)
    print(f"Successfully generated '{file_name}' in the current directory.")

if __name__ == "__main__":
    create_document()
