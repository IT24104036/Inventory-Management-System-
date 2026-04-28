from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# Set margins to 1 inch on all sides
for section in doc.sections:
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

# Add page numbers in footer
def add_page_numbers(doc):
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        para = footer.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run("Page ")
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')
        instrText = OxmlElement('w:instrText')
        instrText.text = 'PAGE'
        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'end')
        run2 = para.add_run()
        run2._r.append(fldChar1)
        run2._r.append(instrText)
        run2._r.append(fldChar2)
        run3 = para.add_run(" of ")
        run3.font.name = "Calibri"
        run3.font.size = Pt(10)
        fldChar3 = OxmlElement('w:fldChar')
        fldChar3.set(qn('w:fldCharType'), 'begin')
        instrText2 = OxmlElement('w:instrText')
        instrText2.text = 'NUMPAGES'
        fldChar4 = OxmlElement('w:fldChar')
        fldChar4.set(qn('w:fldCharType'), 'end')
        run4 = para.add_run()
        run4._r.append(fldChar3)
        run4._r.append(instrText2)
        run4._r.append(fldChar4)

add_page_numbers(doc)

def set_normal_font(run):
    run.font.name = "Calibri"
    run.font.size = Pt(11)

def add_normal_paragraph(doc, text):
    para = doc.add_paragraph()
    run = para.add_run(text)
    set_normal_font(run)
    return para

def add_bullet(doc, text, bold_prefix=None):
    para = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        rb = para.add_run(bold_prefix)
        rb.bold = True
        rb.font.name = "Calibri"
        rb.font.size = Pt(11)
        r = para.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(11)
    else:
        r = para.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(11)
    return para

def shade_cell(cell, fill_hex):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

# ─────────────────────────────────────────
# TITLE PAGE
# ─────────────────────────────────────────
doc.add_paragraph()
doc.add_paragraph()

title_para = doc.add_paragraph()
title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_run = title_para.add_run("Aqsa's Part – Sales Module")
title_run.bold = True
title_run.font.name = "Calibri"
title_run.font.size = Pt(26)
title_run.font.color.rgb = RGBColor(0x1F, 0x39, 0x64)

doc.add_paragraph()

sub_para = doc.add_paragraph()
sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_run = sub_para.add_run("IT2021: AIML Project – Assignment 03 Progress II")
sub_run.bold = True
sub_run.font.name = "Calibri"
sub_run.font.size = Pt(14)
sub_run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)

doc.add_paragraph()

system_para = doc.add_paragraph()
system_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
sys_run = system_para.add_run("System: Invigo FreshGuard")
sys_run.italic = True
sys_run.font.name = "Calibri"
sys_run.font.size = Pt(13)

doc.add_paragraph()
doc.add_paragraph()

for label, value in [
    ("Student:", "Aqsa"),
    ("Responsible Module:", "Sales Module (Sales Recording)"),
    ("Institution:", "Sri Lanka Institute of Information Technology"),
    ("Department:", "Department of IT, Faculty of Computing"),
    ("Year / Semester:", "2nd Year, Semester 2, 2026"),
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rb = p.add_run(f"{label}  ")
    rb.bold = True
    rb.font.name = "Calibri"
    rb.font.size = Pt(12)
    rv = p.add_run(value)
    rv.font.name = "Calibri"
    rv.font.size = Pt(12)

doc.add_page_break()

# ─────────────────────────────────────────
# SECTION 1
# ─────────────────────────────────────────
h1 = doc.add_heading("Section 1: Responsible Components (~75% Implemented)", level=1)
for run in h1.runs:
    run.font.name = "Calibri"

add_normal_paragraph(doc,
    "The following features of the Sales Module at /staff/sales have been implemented and are "
    "functional at approximately 75% completion. Each component is described below.")

doc.add_paragraph()

components = [
    (
        "1. POS Bill Recording Form",
        "Staff members can record multi-product bills through a structured point-of-sale form. "
        "The form supports product selection, quantity input, sale date, optional notes, and optional "
        "customer information. Upon submission, stock is automatically deducted from inventory batches "
        "using the FEFO (First Expired, First Out) principle, ensuring that the soonest-to-expire stock "
        "is consumed first. Each bill can contain multiple line items, making it practical for real "
        "grocery store use."
    ),
    (
        "2. Draft Bill Support",
        "Bills can be saved as drafts without triggering any stock deduction. This is useful when a "
        "transaction is not yet finalised or when data entry needs to be paused. Draft bills are stored "
        "and can be reviewed in the Sales History Table. A 'Finalize' button is available on draft entries, "
        "which, when clicked, processes the bill fully and deducts stock at that point in time."
    ),
    (
        "3. Sales History Table",
        "A comprehensive table displays all sales records. Users can filter by status (All, Active, Draft, "
        "Voided, or Edited), by date range, and by product name using a search input. Sorting options "
        "include Newest First, By Product, and By Quantity. Each row displays the bill ID, product name, "
        "quantity, amount, date, and a colour-coded status badge, giving staff and admins a clear overview "
        "of all recorded transactions."
    ),
    (
        "4. Full Bill Edit (Replace)",
        "Administrators and authorised staff members can perform a full replacement edit on an existing bill. "
        "When an edit is submitted, the original bill lines are voided and inventory is restored, after which "
        "the new line items replace them. Providing an edit reason is mandatory. The system preserves a full "
        "audit trail, recording the editor's identity, the timestamp of the edit, and the stated reason."
    ),
    (
        "5. Void Sale",
        "Any active sale line can be voided by authorised users. Upon voiding, the associated inventory is "
        "restored to the latest available batch. A void reason is required before the action can be confirmed. "
        "Voided records are retained in the database for audit purposes but are excluded from sales reports "
        "and dashboards."
    ),
    (
        "6. Invoice / Bill Detail Modal",
        "Clicking on a bill ID in the Sales History Table opens a styled invoice modal. The modal presents "
        "all line items, unit prices, any discounts applied, the grand total, and customer information if "
        "provided. It also displays the full audit trail for the bill, including who created it, who edited "
        "it, and who voided it — with timestamps and reasons for each action."
    ),
    (
        "7. Duplicate Sale Detection",
        "Before a new bill is saved, the system checks whether the exact same combination of product, quantity, "
        "and customer was already recorded as an active sale on the same day. If a duplicate is detected, a "
        "warning dialog is presented to the user, offering the options to 'Review Entry' or 'Save Anyway', "
        "preventing accidental duplicate recording while preserving user control."
    ),
    (
        "8. Role-Based Access and Time Restriction",
        "The /staff/sales route is protected and requires an active session stored in localStorage. Staff "
        "members can only edit or void their own bills, and only within 2 hours of the bill's creation time. "
        "After that window, the edit and void buttons are disabled with an explanatory tooltip. Administrators "
        "are not subject to the 2-hour restriction and can edit or void any bill at any time."
    ),
    (
        "9. Summary Cards",
        "The top of the Sales Module page displays dashboard-style summary cards showing Today's Bills count "
        "and the Total Units Sold Today. These provide staff with an at-a-glance overview of the day's "
        "activity without needing to scroll through the full history table."
    ),
    (
        "10. Email Receipt",
        "If a customer's email address is provided when recording or updating a bill, the system can optionally "
        "send a receipt email to the customer. This feature is triggered by checking the 'Send receipt' option "
        "in the POS form and is available both on new bill creation and on bill updates."
    ),
]

for title, desc in components:
    h2 = doc.add_heading(title, level=2)
    for run in h2.runs:
        run.font.name = "Calibri"
    add_normal_paragraph(doc, desc)
    doc.add_paragraph()

doc.add_page_break()

# ─────────────────────────────────────────
# SECTION 2
# ─────────────────────────────────────────
h1 = doc.add_heading("Section 2: User Experience – Task Flow, Navigation & Usability", level=1)
for run in h1.runs:
    run.font.name = "Calibri"

add_normal_paragraph(doc,
    "The Sales Module has been designed with usability at its core. The following elements contribute "
    "to a smooth and efficient user experience for grocery store staff.")

doc.add_paragraph()

ux_points = [
    ("Spatial Layout",
     "The POS recording form is positioned on the left side of the page, while the Sales History Table "
     "occupies the right side. This logical spatial separation allows staff to record a new bill and "
     "simultaneously reference recent transactions without switching views or pages."),
    ("Toast Notifications",
     "Every significant action is confirmed by a toast notification displayed at the corner of the screen. "
     "Actions such as successfully recording a sale, updating a bill, voiding a line, or encountering an "
     "error each trigger a contextually appropriate toast, giving immediate and unambiguous feedback."),
    ("Real-Time Stock Display and Line-Total Calculation",
     "As the user types a quantity into the POS form, the available stock for the selected product is "
     "displayed inline in real time. The line total (quantity × unit price) is also calculated and "
     "displayed immediately, giving the user instant visibility into the financial impact of each line "
     "before submission."),
    ("Filters and Sort Controls",
     "Above the Sales History Table, dedicated filter controls allow users to narrow down records by "
     "status, date range, or product name. Sort options enable ordering by newest, by product, or by "
     "quantity. Together, these controls make it efficient to locate specific transactions without "
     "scrolling through a long list."),
    ("Disabled Buttons with Informative Tooltips",
     "Where an action is not permitted — for example, a staff member attempting to edit a bill older "
     "than 2 hours — the corresponding button is visually disabled. Hovering over it reveals an "
     "explanatory tooltip such as 'Staff can only edit within 2 hours', making the restriction "
     "transparent and reducing user frustration."),
    ("Draft Workflow",
     "The draft save option allows staff to enter partial bill data and step away without committing "
     "stock changes. This is particularly useful during busy periods or when product availability needs "
     "to be confirmed before finalisation. The draft state is clearly visible in the history table."),
    ("Clear Status Badges",
     "Every bill in the Sales History Table is accompanied by a colour-coded status badge: Active "
     "bills display a green badge, Draft bills display a yellow badge, and Voided bills display a "
     "grey badge. This visual differentiation allows users to assess the state of transactions "
     "instantly without reading individual field values."),
]

for title, desc in ux_points:
    h2 = doc.add_heading(title, level=2)
    for run in h2.runs:
        run.font.name = "Calibri"
    add_normal_paragraph(doc, desc)
    doc.add_paragraph()

doc.add_page_break()

# ─────────────────────────────────────────
# SECTION 3
# ─────────────────────────────────────────
h1 = doc.add_heading("Section 3: UI Consistency & Standards", level=1)
for run in h1.runs:
    run.font.name = "Calibri"

add_normal_paragraph(doc,
    "The Sales Module adheres to a defined set of UI standards that ensure visual consistency "
    "throughout the interface, making it easier for users to learn and navigate the system.")

doc.add_paragraph()

ui_points = [
    ("Card Styling",
     "All content sections and form containers use consistent rounded-2xl card styling throughout "
     "the module. Labels are styled with font-black weight and tracking-widest uppercase lettering, "
     "creating a uniform and professional appearance across all pages and modals."),
    ("Colour-Coded Status Badges",
     "Status badges for bill states are implemented using a consistent set of Tailwind CSS classes "
     "across all instances in the application. The same colour mapping — green for Active, yellow "
     "for Draft, grey for Void — is applied uniformly, ensuring users encounter no inconsistencies "
     "between pages."),
    ("Theme by Role",
     "The application applies different visual themes depending on the logged-in user's role. The "
     "Staff theme uses warm tones (#F9F5EC background, #4E342E accents) to create a friendly and "
     "approachable environment. The Admin theme uses a dark palette (#0F172A) to convey authority "
     "and differentiate the administrative interface. The theme is set via a role prop passed to "
     "the layout component."),
    ("Form Structure",
     "All forms across the Sales Module follow the same structural pattern: a label is placed above "
     "the input field, the input field itself follows, and any helper text or error message appears "
     "below. This consistent label → input → helper-text structure makes the forms predictable and "
     "accessible."),
    ("Table Headers",
     "Column headers in all tables use consistent font-black weight, uppercase lettering, 10px font "
     "size, and tracking-widest letter spacing. This standardised treatment distinguishes headers "
     "from data rows and maintains the same visual language across all table instances in the module."),
    ("Action Buttons",
     "All action buttons follow a consistent design pattern pairing an icon with a text label. Hover "
     "states are uniform across all buttons of the same type, and destructive actions (void, delete) "
     "consistently use red coloring to signal their impact. This reduces cognitive load for users "
     "who encounter buttons in multiple locations."),
]

for title, desc in ui_points:
    h2 = doc.add_heading(title, level=2)
    for run in h2.runs:
        run.font.name = "Calibri"
    add_normal_paragraph(doc, desc)
    doc.add_paragraph()

doc.add_page_break()

# ─────────────────────────────────────────
# SECTION 4
# ─────────────────────────────────────────
h1 = doc.add_heading("Section 4: Input Validation & Error Handling", level=1)
for run in h1.runs:
    run.font.name = "Calibri"

add_normal_paragraph(doc,
    "Comprehensive validation is implemented at both the frontend and backend levels to ensure "
    "data integrity and prevent invalid or malicious submissions.")

doc.add_paragraph()

# POS Form
h2 = doc.add_heading("POS Form (New Bill)", level=2)
for run in h2.runs:
    run.font.name = "Calibri"

pos_validations = [
    "At least one product must be selected before the form can be submitted.",
    "Quantity must be a positive integer — zero and negative values are rejected.",
    "Quantity cannot exceed the available stock for the selected product; an inline red warning with an alert icon is displayed if this limit is exceeded.",
    "Sale date is a required field and cannot be left blank.",
    "The Submit button remains disabled until all fields pass the isFormValid check, preventing premature submission.",
    "If the 'Send receipt' option is checked, the customer email address is validated against a regex pattern before saving.",
    "All invalid inputs are blocked at the form level before any API call is made.",
]
for v in pos_validations:
    add_bullet(doc, v)

doc.add_paragraph()

# Edit Bill Modal
h2 = doc.add_heading("Edit Bill Modal", level=2)
for run in h2.runs:
    run.font.name = "Calibri"

edit_validations = [
    "At least one valid product line must be present in the edit form.",
    "An edit reason is a mandatory field, marked with a red asterisk; the Save button is disabled until a reason is provided.",
    "Duplicate product IDs within the same bill are detected and blocked with a clear error message.",
    "If the email field is populated and 'Send receipt' is checked, the email format is validated with a regex before saving.",
    "Stock availability is re-validated for each line item at the point of saving, accounting for any stock changes since the modal was opened.",
    "An authorisation re-check is performed at save time: if a staff member's 2-hour edit window has expired by the time they click Save, the action is blocked with a 'Permission Denied' toast notification.",
]
for v in edit_validations:
    add_bullet(doc, v)

doc.add_paragraph()

# Void Dialog
h2 = doc.add_heading("Void Dialog", level=2)
for run in h2.runs:
    run.font.name = "Calibri"

add_normal_paragraph(doc,
    "A void reason is required before a sale can be voided. The Confirm button in the void dialog "
    "remains disabled until the reason field contains a non-empty value.")

doc.add_paragraph()

# Backend Validation
h2 = doc.add_heading("Backend Validation (SaleService.java)", level=2)
for run in h2.runs:
    run.font.name = "Calibri"

backend_validations = [
    "productId is a required field; requests without it are rejected.",
    "quantity must be greater than zero; non-positive values are rejected.",
    "Sufficient stock is checked before deduction; an IllegalStateException is thrown if stock is insufficient.",
    "FEFO deduction consistency is enforced so that the earliest-expiring batch is always consumed first.",
    "Voiding an already-voided sale is blocked at the service level.",
    "Editing a voided or draft sale directly is not permitted; such attempts are rejected with an appropriate error.",
]
for v in backend_validations:
    add_bullet(doc, v)

doc.add_paragraph()

doc.add_page_break()

# ─────────────────────────────────────────
# SECTION 5
# ─────────────────────────────────────────
h1 = doc.add_heading("Section 5: Testing of Responsible Components", level=1)
for run in h1.runs:
    run.font.name = "Calibri"

add_normal_paragraph(doc,
    "Testing has been conducted at both the automated frontend level (using Vitest) and through "
    "structured manual test cases covering all major functionality of the Sales Module.")

doc.add_paragraph()

h2 = doc.add_heading("Frontend Tests (Vitest)", level=2)
for run in h2.runs:
    run.font.name = "Calibri"

test_files = [
    ("authGuard.test.js",
     "12 test cases covering: getSession() returning null when no session exists, "
     "correct parsing of session data, and returning null on invalid JSON input. "
     "ProtectedRoute tests verify that unauthenticated users are blocked, users with "
     "the wrong role are blocked, and users with the correct role are allowed through. "
     "A case-insensitivity test confirms role checking works regardless of letter casing. "
     "GuestRoute tests verify that admins are redirected to /admin, staff to /staff, and "
     "all other non-admin roles to /staff."),
    ("login.test.js",
     "Tests covering the login flow and its validation logic, including handling of "
     "empty fields and incorrect credential formats."),
    ("userManagement.test.js",
     "Tests covering user management operations such as creating, updating, and "
     "listing user accounts."),
    ("staffProfile.test.js",
     "Tests covering staff profile functionality including viewing and updating "
     "profile information."),
]

for filename, desc in test_files:
    h2b = doc.add_heading(filename, level=2)
    for run in h2b.runs:
        run.font.name = "Calibri"
    add_normal_paragraph(doc, desc)
    doc.add_paragraph()

h2 = doc.add_heading("Security-Relevant Test Highlights from authGuard.test.js", level=2)
for run in h2.runs:
    run.font.name = "Calibri"

security_highlights = [
    '"blocks a staff user from accessing admin-only routes" — confirms route-level role enforcement.',
    '"blocks a staff user regardless of role casing" — confirms case-insensitive role comparison.',
    '"returns null when localStorage contains invalid JSON" — confirms defensive JSON parsing to prevent runtime crashes.',
]
for s in security_highlights:
    add_bullet(doc, s)

doc.add_paragraph()

h2 = doc.add_heading("Manual Test Cases for Sales Module", level=2)
for run in h2.runs:
    run.font.name = "Calibri"

add_normal_paragraph(doc,
    "The following manual test cases were conducted and are demonstrable.")

doc.add_paragraph()

# Build table
headers = ["#", "Test Case", "Input", "Expected Outcome", "Result"]
rows = [
    ["1", "Record sale with valid data", "Product + Qty within stock", "Sale recorded, stock deducted", "Pass"],
    ["2", "Submit with empty product", "No product selected", "Submit button disabled", "Pass"],
    ["3", "Quantity exceeds stock", "Qty > available stock", "Red inline warning, submit disabled", "Pass"],
    ["4", "Void with reason", "Void reason typed", "Sale voided, stock restored", "Pass"],
    ["5", "Void without reason", "Empty reason field", "Confirm button disabled", "Pass"],
    ["6", "Edit bill without reason", "Empty edit reason field", "Save button disabled", "Pass"],
    ["7", "Duplicate sale detection", "Same product + qty + customer on same day", "Warning dialog shown", "Pass"],
    ["8", "Staff edits own bill within 2 hrs", "Own bill, < 2 hours old", "Edit allowed", "Pass"],
    ["9", "Staff edits bill after 2 hrs", "Own bill, > 2 hours old", "Edit button disabled", "Pass"],
    ["10", "Staff edits another user's bill", "Different recordedBy value", "Edit button disabled", "Pass"],
    ["11", "Invalid email in POS form", '"notanemail" with Send Receipt checked', "Toast: Invalid Email", "Pass"],
    ["12", "Duplicate product in edit bill", "Same product added twice", "Toast: Duplicate Product", "Pass"],
    ["13", "Edit bill stock re-validation", "Qty > available at save time", "Toast: Insufficient Stock", "Pass"],
    ["14", "Unauthorized route access", "No session in localStorage", "Redirect to /login", "Pass"],
    ["15", "Staff accessing admin route", "Staff session active", "Redirect to /login", "Pass"],
]

col_widths = [Cm(1.0), Cm(4.5), Cm(4.5), Cm(4.8), Cm(1.5)]

table = doc.add_table(rows=1 + len(rows), cols=5)
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

# Header row
hdr_cells = table.rows[0].cells
for i, (cell, header) in enumerate(zip(hdr_cells, headers)):
    cell.width = col_widths[i]
    shade_cell(cell, "2E75B6")
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(header)
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

# Data rows
for r_idx, row_data in enumerate(rows):
    row_cells = table.rows[r_idx + 1].cells
    fill = "EBF3FB" if r_idx % 2 == 0 else "FFFFFF"
    for c_idx, (cell, value) in enumerate(zip(row_cells, row_data)):
        cell.width = col_widths[c_idx]
        shade_cell(cell, fill)
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 4] else WD_ALIGN_PARAGRAPH.LEFT
        run = para.add_run(value)
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        if value == "Pass":
            run.font.color.rgb = RGBColor(0x1F, 0x7A, 0x1F)

doc.add_paragraph()
doc.add_page_break()

# ─────────────────────────────────────────
# SECTION 6
# ─────────────────────────────────────────
h1 = doc.add_heading("Section 6: Backend Security Notes", level=1)
for run in h1.runs:
    run.font.name = "Calibri"

add_normal_paragraph(doc,
    "The following backend security measures are in place within the Invigo FreshGuard system, "
    "relevant to the Sales Module and the overall application.")

doc.add_paragraph()

security_points = [
    ("Password Hashing",
     "Spring Security's SecurityConfig.java uses BCryptPasswordEncoder for all password hashing, "
     "ensuring that passwords are never stored in plain text and are resistant to brute-force attacks."),
    ("Stateless Sessions",
     "The application is configured with a STATELESS session management policy in Spring Security, "
     "meaning no server-side HTTP sessions are maintained. Authentication state is managed client-side "
     "via localStorage."),
    ("API Endpoint Configuration",
     "All sales-related API endpoints are located under /api/sales/* with CORS enabled to allow "
     "cross-origin requests from the frontend application."),
    ("Frontend Role Enforcement",
     "Role-based access control is currently enforced at the frontend route level via the "
     "ProtectedRoute component and permission checks derived from the session stored in localStorage. "
     "The canEditThisSale function enforces time-based access control, restricting staff to a "
     "2-hour window while granting administrators unrestricted access."),
    ("Future Improvement",
     "Server-side re-authentication and authorisation checks on sensitive endpoints (edit, void) "
     "are noted as a future improvement. Currently, these controls are enforced client-side, which "
     "is sufficient for the current prototype stage but should be hardened before production deployment."),
]

for title, desc in security_points:
    h2 = doc.add_heading(title, level=2)
    for run in h2.runs:
        run.font.name = "Calibri"
    add_normal_paragraph(doc, desc)
    doc.add_paragraph()

# Save
output_path = r"C:\Users\ASUS\Downloads\invigo-freshguard-main\invigo-freshguard-main\Aqsa's Part.docx"
doc.save(output_path)
print(f"Saved: {output_path}")